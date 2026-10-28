"""Read the separate ABS business-counts cube (retail, Tables 1–3).

Counts are businesses, not shops. Financial-year entries/exits and the two
survival cohorts remain separate from monthly spending and quarterly sales.
The cube does not publish time-series IDs: this module does not invent ABS IDs.
"""
from datetime import date
import math
from pathlib import Path
import re

from openpyxl import load_workbook


FLOW_LABELS = {
    "Operating at start of financial year": "opening",
    "Entries - Births": "entries_births",
    "Entries - Other": "entries_other",
    "Entries - Total": "entries_total",
    "Exits - Cancellations": "exits_cancellations",
    "Exits - Other": "exits_other",
    "Exits - Total": "exits_total",
    "Operating at end of financial year": "closing",
    "Change": "official_net_change",
    "Percentage change": "official_change_pct",
    "Entry rate": "entry_rate_pct",
    "Exit rate": "exit_rate_pct",
}
RATE_FIELDS = {"official_change_pct", "entry_rate_pct", "exit_rate_pct"}
MISSING_MARKERS = {"", "..", "...", "na", "n/a", "n.a.", "np"}


def _label(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _value(value, field):
    """Preserve known missing markers; reject malformed or nonfinite values."""
    if value is None or (isinstance(value, str) and value.strip().lower() in MISSING_MARKERS):
        return None
    if isinstance(value, bool):
        raise ValueError("Counts field cannot be boolean: " + field)
    try:
        number = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError("Invalid Counts number: " + field) from error
    if not math.isfinite(number):
        raise ValueError("Nonfinite Counts number: " + field)
    is_rate = field in RATE_FIELDS or field == "official_survival_rate_pct"
    if not is_rate and not number.is_integer():
        raise ValueError("Business counts must be integers: " + field)
    if number < 0 and field not in {"official_net_change", "official_change_pct"}:
        raise ValueError("Negative business count or rate: " + field)
    if field in {"entry_rate_pct", "exit_rate_pct", "official_survival_rate_pct"} and number > 100:
        raise ValueError("Counts rate exceeds 100 percent: " + field)
    return number if is_rate else int(number)


def _source(metadata):
    if not isinstance(metadata, dict):
        raise ValueError("Counts source metadata must be a dictionary")
    result = dict(metadata)
    for key in ("url", "release_date", "download_date", "license"):
        if not isinstance(result.get(key), str) or not result[key].strip():
            raise ValueError("Missing Counts source metadata: " + key)
    try:
        released = date.fromisoformat(result["release_date"])
        downloaded = date.fromisoformat(result["download_date"])
    except ValueError as error:
        raise ValueError("Counts source dates must use YYYY-MM-DD") from error
    if released > downloaded:
        raise ValueError("Counts release is after the recorded download date")
    return result, released


def _financial_year(text):
    match = re.fullmatch(r"(\d{4})-(\d{2})", text)
    if not match:
        raise ValueError("Invalid financial-year label: " + text)
    start = int(match[1])
    if int(match[2]) != (start + 1) % 100:
        raise ValueError("Financial year must cover consecutive years")
    return start, date(start + 1, 6, 30)


def _find_header(rows, required):
    matches = [i for i, row in enumerate(rows) if required in [_label(v) for v in row]]
    if len(matches) != 1:
        raise ValueError("Missing or repeated Counts header: " + required)
    return matches[0]


def _check_equal(actual, expected, label):
    if actual is not None and expected is not None and actual != expected:
        raise ValueError("Counts reconciliation failed: " + label)


def _sum(*values):
    return None if any(v is None for v in values) else sum(values)


def _percent(numerator, denominator):
    return None if numerator is None or denominator in (None, 0) else numerator / denominator * 100


def _check_rounded_percent(official, derived, label):
    # Published percentages are rounded to one decimal place.
    if official is not None and derived is not None and abs(official - derived) > 0.050001:
        raise ValueError("Counts rate reconciliation failed: " + label)


def _parse_flows(rows, released):
    header = _find_header(rows, "Operating at start of financial year")
    columns = {}
    for column, text in enumerate(rows[header]):
        label = _label(text)
        if label in FLOW_LABELS:
            field = FLOW_LABELS[label]
            if field in columns:
                raise ValueError("Duplicate Counts flow column: " + field)
            columns[field] = column
    if set(columns) != set(FLOW_LABELS.values()):
        raise ValueError("Counts Table 1 is missing required columns")
    for field, column in columns.items():
        expected = "%" if field in RATE_FIELDS else "no."
        if _label(rows[header + 1][column]) != expected:
            raise ValueError("Unexpected Counts unit for " + field)
    observations = []
    year = None
    for row in rows[header + 2:]:
        first = _label(row[0])
        if re.fullmatch(r"\d{4}-\d{2}", first):
            year = first
        if first.lower() != "retail trade":
            continue
        if year is None:
            raise ValueError("Retail row has no financial-year heading")
        _, end = _financial_year(year)
        if end > released:
            raise ValueError("Counts observation ends after the release date")
        result = {"period": year, "period_end": end.isoformat()}
        result.update({field: _value(row[column], field) for field, column in columns.items()})
        entries, exits = result["entries_total"], result["exits_total"]
        net = None if entries is None or exits is None else entries - exits
        result["derived_net_change"] = net
        result["derived_change_pct"] = _percent(net, result["opening"])
        _check_equal(entries, _sum(result["entries_births"], result["entries_other"]), "entry components")
        _check_equal(exits, _sum(result["exits_cancellations"], result["exits_other"]), "exit components")
        _check_equal(result["official_net_change"], net, "net change")
        _check_equal(result["closing"], _sum(result["opening"], net), "opening + entries - exits")
        _check_rounded_percent(result["official_change_pct"], result["derived_change_pct"], "change")
        _check_rounded_percent(result["entry_rate_pct"], _percent(entries, result["opening"]), "entry")
        _check_rounded_percent(result["exit_rate_pct"], _percent(exits, result["opening"]), "exit")
        result["missing_fields"] = [field for field in columns if result[field] is None]
        result["status"] = "partial" if result["missing_fields"] else "observed"
        observations.append(result)
    if not observations:
        raise ValueError("Counts Table 1 has no retail observations")
    observations.sort(key=lambda r: r["period"])
    years = [_financial_year(r["period"])[0] for r in observations]
    if years != list(range(years[0], years[-1] + 1)):
        raise ValueError("Counts financial years must be unique and continuous")
    for previous, current in zip(observations, observations[1:]):
        _check_equal(current["opening"], previous["closing"], "adjacent financial years")
    units = {field: "percent" if field in RATE_FIELDS else "businesses" for field in columns}
    units.update(derived_net_change="businesses", derived_change_pct="percent")
    return {"table_id": "1", "frequency": "financial_year", "field_units": units,
            "observations": observations}


def _parse_cohort(rows, table, released):
    prefix = "Operating in June " if table == "2" else "Entries in "
    matches = [(i, j, _label(cell)) for i, row in enumerate(rows) for j, cell in enumerate(row)
               if _label(cell).startswith(prefix)]
    if len(matches) != 1:
        raise ValueError("Missing or repeated Counts cohort header")
    header, baseline_column, baseline_label = matches[0]
    suffix = baseline_label[len(prefix):]
    if table == "2":
        if not re.fullmatch(r"\d{4}", suffix):
            raise ValueError("Invalid stock cohort year")
        baseline_end = date(int(suffix), 6, 30)
        cohort_period = suffix + "-06"
    else:
        _, baseline_end = _financial_year(suffix)
        cohort_period = suffix
    if baseline_end > released:
        raise ValueError("Counts cohort baseline is after release")
    if _label(rows[header + 1][baseline_column]) != "no.":
        raise ValueError("Counts cohort baseline unit must be no.")
    retail = [row for row in rows[header + 2:] if _label(row[0]).lower() == "retail trade"]
    if len(retail) != 1:
        raise ValueError("Counts cohort requires one retail row")
    row = retail[0]
    size = _value(row[baseline_column], "cohort_size")
    observations = []
    for column, text in enumerate(rows[header]):
        match = re.fullmatch(r"Survived to June (\d{4})", _label(text))
        if not match:
            continue
        end = date(int(match[1]), 6, 30)
        if not baseline_end < end <= released:
            raise ValueError("Counts survival date outside baseline/release window")
        if column + 1 >= len(row) or _label(rows[header][column + 1]) != "Survival rate":
            raise ValueError("Missing paired survival-rate column")
        if (_label(rows[header + 1][column]), _label(rows[header + 1][column + 1])) != ("no.", "%"):
            raise ValueError("Unexpected survival-count or rate units")
        surviving = _value(row[column], "surviving_businesses")
        rate = _value(row[column + 1], "official_survival_rate_pct")
        if surviving is not None and size is not None and surviving > size:
            raise ValueError("Surviving businesses exceed the fixed cohort")
        derived = _percent(surviving, size)
        _check_rounded_percent(rate, derived, "survival")
        observations.append({"period": end.strftime("%Y-%m"), "period_end": end.isoformat(),
                             "surviving_businesses": surviving, "official_survival_rate_pct": rate,
                             "derived_survival_rate_pct": derived,
                             "status": "partial" if surviving is None or rate is None else "observed"})
    if not observations:
        raise ValueError("No Counts survival observations")
    observations.sort(key=lambda r: r["period"])
    years = [int(r["period"][:4]) for r in observations]
    if years != list(range(years[0], years[-1] + 1)):
        raise ValueError("Counts survival years must be unique and continuous")
    for previous, current in zip(observations, observations[1:]):
        a, b = previous["surviving_businesses"], current["surviving_businesses"]
        if a is not None and b is not None and b > a:
            raise ValueError("Survival count increases within a fixed cohort")
    return {"table_id": table, "cohort_type": "existing_businesses" if table == "2" else "entrants",
            "cohort_period": cohort_period, "cohort_size": size,
            "cohort_frequency": "annual_snapshot" if table == "2" else "financial_year",
            "frequency": "annual_snapshot",
            "field_units": {"cohort_size": "businesses", "surviving_businesses": "businesses",
                            "official_survival_rate_pct": "percent", "derived_survival_rate_pct": "percent"},
            "observations": observations}


def read_business_counts(path, source_metadata):
    """Return the retail annual flows and two distinct cohorts as JSON-safe data.

    source_metadata must contain url, release_date, download_date and license.
    This release's Tables 1–3 have no official Series IDs. The returned table
    and cohort keys are local structure labels, not invented ABS Series IDs.
    """
    source, released = _source(source_metadata)
    path = Path(path)
    if not path.is_file():
        raise ValueError("Missing ABS Counts workbook: " + str(path))
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        tables = {}
        for number in ("1", "2", "3"):
            name = "Table " + number
            if name not in workbook.sheetnames:
                raise ValueError("Missing ABS Counts worksheet: " + name)
            tables[number] = list(workbook[name].iter_rows(values_only=True))
        flows = _parse_flows(tables["1"], released)
        cohorts = [_parse_cohort(tables[n], n, released) for n in ("2", "3")]
    finally:
        workbook.close()
    # The cohorts should reconcile to the corresponding baseline in Table 1.
    by_period = {r["period"]: r for r in flows["observations"]}
    for cohort in cohorts:
        if cohort["cohort_type"] == "entrants":
            baseline = by_period.get(cohort["cohort_period"])
            field = "entries_total"
        else:
            year = int(cohort["cohort_period"][:4])
            baseline = by_period.get(f"{year}-{(year + 1) % 100:02d}")
            field = "opening"
        if baseline:
            _check_equal(cohort["cohort_size"], baseline[field], "cohort vs Table 1 baseline")
    return {"schema_version": 1, "dataset_id": "abs_business_counts", "source": source,
            "industry": "Retail trade", "classification": "ANZSIC", "geography": "Australia",
            "adjustment": "Original", "price_basis": "not_applicable", "official_series_ids": None,
            "annual_flows": flows, "survival_cohorts": cohorts,
            "quality": {"annual_flow_rows": len(flows["observations"]),
                        "survival_rows": sum(len(c["observations"]) for c in cohorts),
                        "missing_flow_values": sum(len(r["missing_fields"]) for r in flows["observations"]),
                        "missing_survival_values": sum(sum(r[k] is None for k in
                            ("surviving_businesses", "official_survival_rate_pct"))
                            for c in cohorts for r in c["observations"]),
                        "reconciliation": "checked_where_inputs_present"},
            "notes": ["Business counts are not store counts.",
                      "Entries include births and other entries; exits are not a bankruptcy measure.",
                      "Existing-business and entrant survival rates have different fixed cohorts.",
                      "Missing observations remain null; derived percentages use unrounded counts.",
                      "No official Series ID is supplied in Data Cube 1 Tables 1–3."]}
